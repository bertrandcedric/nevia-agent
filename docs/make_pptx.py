"""docs/presentation.md as a PowerPoint deck in Aiven's brand template, docs/presentation.pptx.

The text (titles, bullets, the table, the speaker's notes, the animations' captions) comes from the Markdown;
the diagrams are drawn here as native shapes, in the brand palette, and the two "step by step" animations are
a slide per step, as in the Markdown, so that going back a step is going back a slide. Run with python-pptx:

    uv run --with python-pptx docs/make_pptx.py [TEMPLATE.pptx]
"""

import re
import sys
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DOCS = Path(__file__).resolve().parent
TEMPLATE = Path.home() / "Documents/MASTER Aiven Brand 2.0 Slides Template.pptx"
OUT = DOCS / "presentation.pptx"

# The brand palette (the template's theme).
DARK, WHITE, GREY, LINE = "05080F", "FFFFFF", "919294", "444547"
GREEN, TEAL, PURPLE, PINK, ORANGE = "5FFA74", "2ED0CD", "6F64FF", "DF56F2", "FF965E"
TEXT, MONO = "Inter", "Geist Mono Medium"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

COVER, CONTENT, CLOSING = "TITLE_2_2_1_1", "TITLE_AND_BODY_1_4_1_1_1", "BLANK_1_1_1"
EYEBROWS = [
    "", "THE PROBLEM", "NEVIA", "NEVIA", "THE MIGRATION", "THE DATA · M2P", "THE CODE · ACC", "VALIDATE, CUT OVER",
    "HOW IT RUNS", "WHAT NEXT",
]  # fmt: skip


# ---------------------------------------------------------------------- the Markdown


@dataclass
class Slide:
    title: str
    subtitle: str
    bullets: list[str]
    table: list[list[str]]
    notes: str
    animation: tuple[str, int] | None  # (m2p or acc, frame)
    caption: str
    after: str  # text after the diagram (a frame's caption, or a sentence under a table)


def read(md: Path) -> list[Slide]:
    text = md.read_text()
    text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)  # the front matter
    text = re.sub(r"\A\s*<!--.*?-->\s*", "", text, flags=re.S)  # the deck's own comment
    slides = []
    for chunk in re.split(r"\n---\n", text):
        anim = re.search(r"<!-- animation: (\w+), frame (\d+)/\d+ -->", chunk)
        notes = " ".join(n.strip() for n in re.findall(r"<!--(.*?)-->", chunk, re.S) if "animation:" not in n)
        body = re.sub(r"<!--.*?-->", "", chunk, flags=re.S)
        body = re.sub(r"```.*?```", "", body, flags=re.S)
        title = (re.search(r"^# (.+)$", body, re.M) or [""])[0].removeprefix("# ")
        subtitle = (re.search(r"^### (.+)$", body, re.M) or [""])[0].removeprefix("### ")
        bullets, table, after, item = [], [], [], None
        for line in body.splitlines():
            if line.startswith("- "):
                item = line[2:]
                bullets.append(item)
            elif line.startswith("  ") and line.strip() and bullets:
                bullets[-1] += " " + line.strip()
            elif line.startswith("|"):
                cells = [c.strip() for c in line.strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    table.append(cells)
            elif line.strip() and not line.startswith("#"):
                after.append(line.strip())
        caption = " ".join(after) if anim else ""
        slides.append(Slide(title, subtitle, bullets, table, " ".join(notes.split()),
                            (anim[1], int(anim[2])) if anim else None, caption, " ".join(after)))  # fmt: skip
    return slides


# ---------------------------------------------------------------------- text


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def runs(paragraph, text: str, size: float, color: str = WHITE, *, bold: bool = False) -> None:
    """`text` into `paragraph`, its **bold** and `code` as such (code in the mono font, green)."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        r = paragraph.add_run()
        code = part.startswith("`")
        r.text = part.strip("*`") if part.startswith(("**", "`")) else part
        f = r.font
        f.size, f.color.rgb = Pt(size), rgb(GREEN if code else color)
        f.bold = bold or part.startswith("**")
        f.name = MONO if code else TEXT


def fill_text(frame, lines: list[str], size: float, color: str = WHITE, *, space: float = 6, bullet: bool = False):  # noqa: PLR0913
    frame.clear()
    frame.word_wrap = True
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        if bullet:
            r = p.add_run()
            r.text = "▸  "
            r.font.size, r.font.color.rgb, r.font.name = Pt(size), rgb(GREEN), TEXT
        runs(p, line, size, color)
        p.space_after = Pt(space)


def textbox(  # noqa: PLR0913, PLR0917 - where it goes, and what
    slide, x, y, w, h, lines, size, color=WHITE, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None
):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    fill_text(tf, lines, size, color, space=2)
    for p in tf.paragraphs:
        p.alignment = align
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = rgb(fill)
    return tb


# ---------------------------------------------------------------------- diagrams

# Columns and rows of the diagrams' grid, in inches.
COLS = [0.35, 2.25, 4.15, 6.05, 7.95]
ROWS = [1.28, 2.5, 3.72]
W, H = 1.45, 0.72

# A box's kind is its outline's colour, named in the legend; every box is a rounded rectangle, its text centred.
KIND = {"you": WHITE, "computer": TEAL, "db": GREEN, "ai": PURPLE, "check": PINK, "result": ORANGE}
LEGEND = [("you", "Your machine"), ("computer", "Compute"), ("db", "Data"), ("ai", "AI"), ("result", "Result")]


@dataclass
class Node:
    id: str
    name: str
    detail: str
    kind: str
    col: int
    row: int
    frame: int = 0

    @property
    def box(self) -> tuple[float, float, float, float]:
        return COLS[self.col], ROWS[self.row], W, H


@dataclass
class Edge:
    a: str
    b: str
    label: str = ""
    frame: int = 0
    dashed: bool = False


# The animations, as timelines: one column per step, from left to right, the steps on an axis under them. Boxes
# are outlined in their kind's colour (the legend, top right), the same at every step; the step shown has its
# arrows and its dot on the axis in green. Nodes: id: name, detail, kind, x, y, w, frame. Edges: (from, to, label,
# frame, how), how being line, curve (an elbow), branch (a dashed one), down or down-dashed (a bottom to a top),
# bus (dashed: along the row of `from`, then up into `to`: a new branch), drop (along the row of `from`, then down
# into `to`), rise or fall (straight up or down where the two boxes overlap), under (a bottom to a bottom, below
# the rows). Results: id: word, color, and the frame it shows in when not the node's.
def colx(i: float) -> float:
    """Where column `i` starts; a half one sits between two columns."""
    return 0.25 + 1.4 * i


COL = [colx(i) for i in range(7)]
ROW = [1.3, 2.2, 3.1]
CARD_W, CARD_H = 1.2, 0.55

M2P_NODES = {
    "you": ("Your machine", "agent m2p", "you", COL[0], ROW[0], CARD_W, 1),
    "agent": ("m2p-agent", "the AI agent, mongo2pg", "computer", COL[0], ROW[1], CARD_W, 1),
    "mongo": ("m2p-mongo", "MongoDB and its data", "db", COL[1], ROW[1], CARD_W, 2),
    "pg": ("m2p-pg", "the target", "db", COL[1], ROW[2], CARD_W, 2),
    "infer": ("mongo2pg infer", "a model, warnings", "ai", COL[2], ROW[1], CARD_W, 3),
    "try1": ("m2p-try1", "a new branch", "db", COL[3], ROW[1], CARD_W, 4),
    "fix1": (
        "Analyse, fix",
        "SQL, or the model",
        "ai",
        colx(3.5) + 0.1,
        ROW[0],
        1.0,
        5,
    ),  # narrower: clear of try2's rise
    "try2": ("m2p-try2", "a new branch", "db", COL[4], ROW[1], CARD_W, 6),
    "fix2": ("Analyse, fix", "SQL, or the model", "ai", colx(4.5) + 0.1, ROW[0], 1.0, 7),
    "try3": ("m2p-try3", "a new branch", "db", COL[5], ROW[1], CARD_W, 8),
    "finish": ("final.sql", "DONE.md, the tries", "result", COL[6], ROW[1], CARD_W, 9),
}
M2P_EDGES = [
    ("you", "agent", "", 1, "down"),
    ("agent", "mongo", "", 2, "line"),
    ("agent", "pg", "", 2, "curve"),
    ("mongo", "infer", "", 3, "line"),
    ("infer", "try1", "", 4, "line"),
    ("pg", "try1", "branch.sh 1", 4, "bus"),
    ("try1", "fix1", "", 5, "rise"),  # the AI reads the failed report, and fixes
    ("fix1", "try2", "", 6, "fall"),
    ("pg", "try2", "branch.sh 2", 6, "bus"),
    ("try2", "fix2", "", 7, "rise"),
    ("fix2", "try3", "", 8, "fall"),
    ("pg", "try3", "branch.sh 3", 8, "bus"),
    ("try3", "finish", "", 9, "line"),
    ("finish", "pg", "", 9, "under"),  # the passing load goes into the target
]
M2P_RESULTS = {"try1": ("FAILS", "F2545B"), "try2": ("FAILS", "F2545B"), "try3": ("PASSES", GREEN),
               "pg": ("LOADED", GREEN, 9)}  # fmt: skip
M2P_AXIS = [("start", 0), ("source, target", 1), ("analyse", 2), ("try 1", 3), ("fix", 3.5), ("try 2", 4),
            ("fix", 4.5), ("try 3", 5), ("done", 6)]  # fmt: skip
M2P = (M2P_NODES, M2P_EDGES, M2P_RESULTS, 8, M2P_AXIS)

ACC_NODES = {
    "you": ("Your machine", "agent acc", "you", COL[0], ROW[0], CARD_W, 1),
    "agent": ("acc-agent", "runs the run", "computer", COL[0], ROW[1], CARD_W, 1),
    "gh": ("GitHub", "clone of the app", "db", COL[1], ROW[0], CARD_W, 2),
    "app": ("acc-app", "the app, from main", "computer", COL[1], ROW[1], CARD_W, 2),
    "mongo": ("acc-mongo", "its MongoDB, seeded", "db", COL[1], ROW[2], CARD_W, 2),
    "infer": ("mongo2pg infer", "without the values", "ai", COL[2], ROW[1], CARD_W, 3),
    "ai": ("acc-ai", "the AI: code, tests", "ai", COL[3], ROW[1], CARD_W, 4),
    "build": ("Build", "the app and its tests", "computer", COL[4], ROW[1], CARD_W, 5),
    "sb": ("acc-app-RUN", "from a checkpoint", "computer", COL[5], ROW[1], CARD_W, 6),
    "sbdb": ("acc-pg-RUN", "the AI's postgres-init", "db", COL[5], ROW[2], CARD_W, 6),
    "pr": ("Pull request", "from agent/RUN", "result", COL[6], ROW[1], CARD_W, 7),
}
ACC_EDGES = [
    ("you", "agent", "", 1, "down"),
    ("agent", "gh", "", 2, "curve"),
    ("agent", "app", "", 2, "line"),
    ("app", "mongo", "", 2, "down-dashed"),
    ("mongo", "infer", "", 3, "curve"),
    ("infer", "ai", "", 4, "line"),  # the AI works from mongo2pg's analysis
    ("gh", "ai", "", 4, "drop"),  # and from the clone
    ("ai", "build", "", 5, "line"),
    ("build", "sb", "", 6, "line"),
    ("sb", "sbdb", "", 6, "down"),
    ("sb", "pr", "", 7, "line"),
]
ACC_AXIS = [(label, i) for i, label in enumerate(["start", "a fresh stack", "analyse", "change", "build", "sandbox",
                                                  "pull request"])]  # fmt: skip
# At the end, both apps are up on Nevia to try side by side: today's on MongoDB, the change on Postgres.
ACC_RESULTS = {"app": ("LIVE", GREEN, 7), "sb": ("LIVE", GREEN, 7)}
ACC = (ACC_NODES, ACC_EDGES, ACC_RESULTS, 8, ACC_AXIS)
OUTLINE, DIM = "3A3D44", "6E7078"

RUNS = (
    [
        Node("you", "Your machine", "agent m2p | acc", "you", 0, 1),
        Node("agent", "Agent computer", "this code, an AI agent", "computer", 2, 1),
        Node("rest", "Everything else", "MongoDB, Postgres, branches, sandboxes, PRs", "db", 4, 1),
    ],
    [Edge("you", "agent", "computer, key, secrets"), Edge("agent", "rest", "makes and deletes")],
)


def box(slide, x, y, w, h, name, detail, kind, *, size=10.0):  # noqa: PLR0913, PLR0917
    """A rounded box, dark, its outline its kind's colour, its name over its detail, centred."""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.13
    s.shadow.inherit = False
    s.line.color.rgb = rgb(KIND[kind])
    s.line.width = Pt(1)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(DARK)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, name, size, bold=True)
    q = tf.add_paragraph()
    q.alignment = PP_ALIGN.CENTER
    q.space_before = Pt(1.5)
    runs(q, detail, round(size * 0.8, 1), GREY)
    return s


def node_shape(slide, n: Node, box_: tuple[float, float, float, float] | None = None):
    return box(slide, *(box_ or n.box), n.name, n.detail, n.kind)


def anchors(a: Node, b: Node) -> tuple[float, float, float, float]:
    """Where an edge leaves `a` and reaches `b`: their facing sides' middles."""
    ax, ay, aw, ah = a.box
    bx, by, bw, bh = b.box
    acx, acy, bcx, bcy = ax + aw / 2, ay + ah / 2, bx + bw / 2, by + bh / 2
    if a.row == b.row or abs(bcx - acx) > abs(bcy - acy) * 1.6:
        x1 = ax + aw if bcx > acx else ax
        x2 = bx if bcx > acx else bx + bw
        return x1, acy, x2, bcy
    y1 = ay + ah if bcy > acy else ay
    y2 = by if bcy > acy else by + bh
    return acx, y1, bcx, y2


def arrow(slide, x1, y1, x2, y2, color=GREY, width=1.0, *, dashed=False, head=True):  # noqa: PLR0913, PLR0917
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        ln.append(etree.fromstring(f'<a:prstDash xmlns:a="{A}" val="dash"/>'))
    if head:
        ln.append(etree.fromstring(f'<a:tailEnd xmlns:a="{A}" type="triangle" w="med" len="med"/>'))
    return c


def edge_shapes(slide, e: Edge, nodes: dict[str, Node]) -> list:
    x1, y1, x2, y2 = anchors(nodes[e.a], nodes[e.b])
    out = [arrow(slide, x1, y1, x2, y2, dashed=e.dashed)]
    width = 0.065 * len(e.label) + 0.12  # at 7 pt
    if e.label and (y1 != y2 or width < abs(x2 - x1) - 0.04):  # only where it fits: the captions say the rest
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2 - (0.17 if y1 == y2 else 0)
        out.append(textbox(slide, mx - width / 2, my - 0.11, width, 0.22, [e.label], 7, GREY, align=PP_ALIGN.CENTER,
                           anchor=MSO_ANCHOR.MIDDLE, fill=DARK))  # fmt: skip
    return out


def triangle(slide) -> None:
    """The migration: the app at the top, MongoDB and Postgres at the base, and where m2p and acc act."""
    node_shape(slide, Node("app", "The app", "its code, its API", "computer", 0, 0), (2.65, 1.3, 1.7, 0.62))
    node_shape(slide, Node("mongo", "MongoDB", "the source", "db", 0, 0), (0.6, 3.55, 1.5, 0.8))
    node_shape(slide, Node("pg", "PostgreSQL", "the target", "db", 0, 0), (4.9, 3.55, 1.5, 0.8))
    # the sides: the app's link to its database, before and after
    arrow(slide, 2.95, 1.92, 1.35, 3.55, dashed=True, head=False)
    arrow(slide, 4.05, 1.92, 5.65, 3.55)
    textbox(slide, 0.6, 2.45, 1.4, 0.25, ["today"], 9, GREY, align=PP_ALIGN.RIGHT)
    textbox(slide, 5.0, 2.45, 1.4, 0.25, ["after"], 9, GREY)
    # mongo2pg's assessment of MongoDB, in the middle: the input both agents start from
    node_shape(slide, Node("assess", "Assessment", "mongo2pg: types, a model", "ai", 0, 0), (2.7, 3.12, 1.6, 0.46))
    arrow(slide, 2.0, 3.62, 2.7, 3.35, dashed=True)
    arrow(slide, 3.5, 3.12, 3.5, 2.9, dashed=True)
    arrow(slide, 3.5, 3.58, 3.5, 3.9, dashed=True)
    # acc moves the app's link from one side to the other, its words above; m2p the data along the base, below
    for y, color, name, detail, x1, x2, ty in (
        (2.85, PURPLE, "acc · moves the code", "JDBC, tested in a sandbox", 2.4, 4.6, 2.22),
        (3.95, PINK, "m2p · moves the data", "tries on branches, checked row by row", 2.2, 4.8, 4.0),
    ):
        arrow(slide, x1, y, x2, y, color, 3)
        textbox(slide, 1.9, ty, 3.2, 0.26, [f"**{name}**"], 11, color, align=PP_ALIGN.CENTER)
        textbox(slide, 1.9, ty + 0.26, 3.2, 0.24, [detail], 8.5, GREY, align=PP_ALIGN.CENTER)


def legend(slide) -> None:
    """The kinds' colours, top right: a dot and a word each."""
    widths = [0.2 + 0.042 * len(word) for _, word in LEGEND]
    x = 9.43 - sum(widths)
    for (kind, word), w in zip(LEGEND, widths, strict=True):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(0.47), Inches(0.07), Inches(0.07))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(KIND[kind])
        dot.line.fill.background()
        dot.shadow.inherit = False
        textbox(slide, x + 0.08, 0.42, w - 0.08, 0.16, [word], 6.2, GREY)
        x += w


def route(how: str, a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> list:
    """An edge's legs, (x1, y1, x2, y2) each, from box `a` to box `b`; the last one gets the arrowhead."""
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    if how == "under":
        x1, x2, y = ax + aw / 2, bx + bw / 2, max(ay + ah, by + bh) + 0.15
        return [(x1, ay + ah, x1, y), (x1, y, x2, y), (x2, y, x2, by + bh)]
    if how in ("rise", "fall"):  # straight up or down, where the two boxes overlap (a try and its fix, side by side)
        x = (max(ax, bx) + min(ax + aw, bx + bw)) / 2
        return [(x, ay, x, by + bh)] if how == "rise" else [(x, ay + ah, x, by)]
    if how in ("bus", "drop"):
        x2, y = bx + bw / 2 + (0.2 if how == "drop" else 0), ay + ah / 2
        return [(ax + aw, y, x2, y), (x2, y, x2, by + bh if how == "bus" else by)]
    if how.startswith("down"):
        return [(ax + 0.25, ay + ah, bx + 0.25, by)]
    if how in ("curve", "branch"):  # an elbow, halfway between the two
        x1, y1, x2, y2 = ax + aw, ay + ah / 2, bx, by + bh / 2
        mx = (x1 + x2) / 2
        return [(x1, y1, mx, y1), (mx, y1, mx, y2), (mx, y2, x2, y2)]
    return [(ax + aw, ay + ah / 2, bx, by + bh / 2)]


def timeline(slide, spec) -> tuple[dict[int, list], dict[int, list]]:
    """The shapes by the frame in which they appear, and the highlights (green) shown during that frame only: its
    arrows and its dot on the axis. A box keeps its colour from one step to the next."""
    nodes, edges, results, size, axis = spec
    shown: dict[int, list] = {}
    lit: dict[int, list] = {}
    legend(slide)
    boxes = {k: (x, y, w, CARD_H) for k, (_, _, _, x, y, w, _) in nodes.items()}
    for a, b, label, frame, how in edges:
        ax, ay, aw, ah = boxes[a]
        bx, by, bw, bh = boxes[b]
        legs = route(how, boxes[a], boxes[b])
        dashed = how in ("branch", "down-dashed", "bus")
        for color, into in ((DIM, shown), (GREEN, lit)):
            for n, pts in enumerate(legs, 1):
                into.setdefault(frame, []).append(arrow(slide, *pts, color, dashed=dashed, head=n == len(legs)))
        if label:  # a new branch's by its tap; a line's above it
            mono = how == "bus"
            lx, ly = (bx + bw / 2 + 0.04, by + bh + 0.1) if mono else ((ax + aw + bx) / 2 - 0.4, ay + ah / 2 - 0.26)
            align = PP_ALIGN.LEFT if mono else PP_ALIGN.CENTER
            tb = textbox(slide, lx, ly, 0.8, 0.2, [f"`{label}`" if mono else label], 6.5, GREY, align=align)
            shown[frame].append(tb)
    for k, (name, detail, kind, *_, frame) in nodes.items():
        x, y, w, h = boxes[k]
        shown.setdefault(frame, []).append(box(slide, x, y, w, h, name, detail, kind, size=size))
        word, color, *later = results.get(k, ("", ""))
        if word:
            at = later[0] if later else frame
            badge = textbox(
                slide,
                x + w - 0.95,
                y + 0.02,
                0.89,
                0.16,
                [f"**{word}**"],
                6,
                color,
                align=PP_ALIGN.RIGHT,
                anchor=MSO_ANCHOR.TOP,
            )  # in the top right corner, over the name
            shown.setdefault(at, []).append(badge)
    steps(slide, axis, shown, lit)
    return shown, lit


def steps(slide, axis: list[str], shown: dict[int, list], lit: dict[int, list]) -> None:
    """The time axis under a timeline: a dot per step, filled once reached, green while shown."""
    if not axis:
        return
    y, r = 4.0, 0.06
    xs = [colx(i) + CARD_W / 2 for _, i in axis]
    arrow(slide, xs[0], y, xs[-1] + 0.35, y, OUTLINE, 1)
    for n, (x, (label, _)) in enumerate(zip(xs, axis, strict=True), 1):
        for color, into in ((OUTLINE, None), (DIM, shown), (GREEN, lit)):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r))
            dot.fill.solid()
            dot.fill.fore_color.rgb = rgb(DARK if into is None else color)
            dot.line.color.rgb = rgb(color)
            dot.shadow.inherit = False
            if into is not None:
                into.setdefault(n, []).append(dot)
        textbox(slide, x - 0.6, y + 0.1, 1.2, 0.22, [f"{n} · {label}"], 7, DIM, align=PP_ALIGN.CENTER)
        tb = textbox(slide, x - 0.6, y + 0.1, 1.2, 0.22, [f"**{n} · {label}**"], 7, GREEN, align=PP_ALIGN.CENTER,
                     fill=DARK)  # fmt: skip
        lit[n].append(tb)


def diagram(slide, spec) -> dict[int, list]:
    """The diagram's shapes, by the frame in which they appear (0: at once)."""
    nodes, edges = spec
    by_id = {n.id: n for n in nodes}
    frames: dict[int, list] = {}
    for e in edges:  # under the nodes
        frames.setdefault(e.frame, []).extend(edge_shapes(slide, e, by_id))
    for n in nodes:
        frames.setdefault(n.frame, []).append(node_shape(slide, n))
    return frames


# ---------------------------------------------------------------------- slides


def layout(prs, name: str):
    return next(lay for m in prs.slide_masters for lay in m.slide_layouts if lay.name == name)


def ph(slide, idx: int):
    return next(s for s in slide.placeholders if s.placeholder_format.idx == idx)


def drop(slide, idx: int) -> None:
    s = ph(slide, idx)
    s._element.getparent().remove(s._element)


def place(shape, top: float, height: float) -> None:
    """A placeholder lower on its slide: all four, or the ones not set fall back to 0."""
    shape.left, shape.top, shape.width, shape.height = Inches(0.6), Inches(top), Inches(8.81), Inches(height)


def notes(slide, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def content(prs, eyebrow: str, title: str, footer: str = "From MongoDB to PostgreSQL, with AI agents on Nevia"):
    s = prs.slides.add_slide(layout(prs, CONTENT))
    ph(s, 3).text_frame.text = eyebrow
    ph(s, 0).text_frame.text = title
    ph(s, 2).text_frame.text = footer
    return s


def bullets_into(placeholder, items: list[str], size: float = 13) -> None:
    fill_text(placeholder.text_frame, items, size, space=9, bullet=True)
    for p in placeholder.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        pPr.insert(0, etree.fromstring('<a:buNone xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))


def table_into(slide, rows: list[list[str]], y: float) -> None:
    n, m = len(rows), len(rows[0])
    t = slide.shapes.add_table(n, m, Inches(0.59), Inches(y), Inches(8.8), Inches(0.42 * n)).table
    t.columns[0].width = Inches(2.0)
    t.columns[1].width = Inches(6.8)
    tbl = t._tbl
    tbl.tblPr.set("firstRow", "1")
    tbl.tblPr.set("bandRow", "0")
    style = tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tbl.tblPr.remove(style)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(DARK)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.clear()
            runs(tf.paragraphs[0], value, 10 if r else 9, GREEN if r == 0 else WHITE, bold=r == 0 or c == 0)
            tcPr = cell._tc.get_or_add_tcPr()
            border = etree.SubElement(tcPr, qn("a:lnB"), w="9525")
            solid = etree.SubElement(border, qn("a:solidFill"))
            etree.SubElement(solid, qn("a:srgbClr"), val=LINE)


def prune(prs) -> None:
    """Only the template's layouts and masters the deck uses: Keynote drops the others' picture placeholders, and
    warns of it as "unsupported media"."""
    used = {sl.slide_layout.part for sl in prs.slides}
    for master in prs.slide_masters:
        for lay in list(master.slide_layouts):
            if lay.part not in used:
                master.slide_layouts.remove(lay)
    ids = prs.slide_masters._sldMasterIdLst
    for sid, master in zip(list(ids), list(prs.slide_masters), strict=True):
        if not len(master.slide_layouts):
            prs.part.drop_rel(sid.rId)
            ids.remove(sid)


def build(template: Path) -> None:
    prs = Presentation(str(template))
    ids = prs.slides._sldIdLst
    for sid in list(ids):  # the template's samples go; its masters and layouts stay
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)

    md = read(DOCS / "presentation.md")
    main = [s for s in md if not s.animation]
    frames = {k: [s for s in md if s.animation and s.animation[0] == k] for k in ("m2p", "acc")}

    # 1. The cover.
    c = main[0]
    s = prs.slides.add_slide(layout(prs, COVER))
    ph(s, 0).text_frame.text = c.title.upper()
    ph(s, 1).text_frame.text = c.subtitle
    ph(s, 2).text_frame.text = "NEVIA · TWO AI AGENTS · M2P AND ACC"
    notes(s, " ".join(c.bullets) + " " + c.notes)

    for i, sl in enumerate(main[1:], 1):
        s = content(prs, EYEBROWS[i], sl.title)
        body = ph(s, 1)
        if sl.table:  # Nevia's concepts
            drop(s, 1)
            table_into(s, sl.table, 1.35)
            textbox(s, 0.59, 4.7, 8.8, 0.3, [sl.after], 11, GREY)
        elif sl.title.startswith("The migration"):
            triangle(s)
            body.left, body.top, body.width, body.height = Inches(6.8), Inches(1.3), Inches(2.8), Inches(3.2)
            body.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            bullets_into(body, sl.bullets, 10.5)
        elif sl.title.startswith("How it runs"):
            diagram(s, RUNS)
            for shape in s.shapes:  # this diagram on the grid's first row
                if not shape.is_placeholder:
                    shape.top -= Inches(1.0)
            place(body, 2.75, 2.3)
            bullets_into(body, sl.bullets, 12)
        else:
            bullets_into(body, sl.bullets)
        notes(s, sl.notes)
        for agent, spec in (("m2p", M2P), ("acc", ACC)):
            if sl.title.endswith(f": {agent}"):
                animated(prs, agent, spec, frames[agent])

    s = prs.slides.add_slide(layout(prs, CLOSING))
    ph(s, 0).text_frame.text = "THANK YOU"
    notes(s, "agent m2p setup && agent m2p run; agent acc setup && agent acc run migrate")
    prune(prs)
    prs.save(str(OUT))
    print(f"{OUT}: {len(prs.slides)} slides")  # noqa: T201 - a script


def animated(prs, agent: str, spec, frames: list[Slide]) -> None:
    """The agent's steps, one slide each, a click apart: what a step adds is drawn from its slide on, green on its
    own. Going back a step is going back a slide, in any app, a PDF included."""
    for n, f in enumerate(frames, 1):
        s = content(prs, f"{agent.upper()} · STEP BY STEP", f"How {agent} works, step by step")
        drop(s, 1)
        shown, lit = timeline(s, spec)
        later = [x for k, xs in shown.items() if k > n for x in xs] + [x for k, xs in lit.items() if k != n for x in xs]
        for shape in later:
            shape._element.getparent().remove(shape._element)
        title, _, rest = f.caption.partition("** ")
        textbox(s, 0.59, 4.6, 8.8, 0.62, [f"**{n}/{len(frames)} · {title.strip('*')}** {rest}"], 11)
        notes(s, f.caption.replace("**", ""))


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else TEMPLATE)
